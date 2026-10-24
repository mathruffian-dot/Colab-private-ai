"""開機自我檢查：確認文件工具「能用」，不只是「裝了」。

在暫存資料夾產生測試檔（不碰工作區、不碰學生資料）：
- 中文圖表：用 matplotlib／seaborn 畫圖，攔截「字型缺字」警告（中文變方塊時一定會出現）；
  graphviz 用 fc-match 確認找得到中文字型。
- OCR：把一行中文畫成乾淨的圖片，用繁中 OCR 單行模式讀回來。
- 產生含表格與原生圖表的 Word／Excel／PowerPoint → 轉成 PDF → 從 PDF 讀回中文；
  另檢查 xlsxwriter、Markdown 表格、plotly 離線 HTML、合併列印、PDF／壓縮檔／Office 加密、
  掃描 PDF 轉文字、公文字型對應、簡轉繁與斷詞、影音、QR code／注音／文字雲、語音合成與辨識、pandoc。
任何一項失敗只顯示警告，不會擋住 AI 啟動。
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile
import warnings

TEXT = "學生成績統計"
SCORES = [("國文", 80), ("數學", 72), ("英文", 90)]
CJK_FILE = "NotoSansCJK"  # fonts-noto-cjk 的字型檔名
results = []


def check(name):
    def wrap(fn):
        try:
            detail = fn()
            results.append((True, name, detail or ""))
        except Exception as e:  # 自我檢查要把每一項都跑完，不因單項失敗中斷
            results.append((False, name, f"{type(e).__name__}: {e}"))
        return fn
    return wrap


def compact(s):
    return "".join(s.split())


def save_without_missing_glyphs(fig, path):
    """存圖時攔截 matplotlib 的「Glyph … missing from font」警告：有缺字就代表中文會變方塊。"""
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        fig.savefig(path)
    missing = [str(w.message) for w in caught if "missing from font" in str(w.message)]
    if missing:
        raise RuntimeError(f"中文缺字（會變方塊）：{missing[0][:80]}")


work = tempfile.mkdtemp(prefix="tools_selfcheck_")
os.chdir(work)


@check("內建套件（pandas、numpy、matplotlib、Pillow、OpenCV）")
def _():
    import cv2, matplotlib, numpy, pandas, PIL  # noqa: F401
    return f"pandas {pandas.__version__}、OpenCV {cv2.__version__}"


@check("matplotlib 中文圖表（不缺字）")
def _():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    font = font_manager.findfont(font_manager.FontProperties(family=["sans-serif"]))
    if CJK_FILE not in font:
        raise RuntimeError(f"預設字型不是中文字型：{os.path.basename(font)}")
    fig, ax = plt.subplots(figsize=(6, 3), dpi=150)
    ax.bar([s for s, _ in SCORES], [v for _, v in SCORES])
    ax.set_title(TEXT, fontsize=28)
    save_without_missing_glyphs(fig, "chart.png")
    plt.close(fig)
    return f"字型：{os.path.basename(font)}"


@check("seaborn 中文統計圖（不缺字）")
def _():
    import matplotlib.pyplot as plt
    import seaborn as sns
    sns.set_theme(font=plt.rcParams["font.sans-serif"][0])  # 與 AGENTS.md 寫給 AI 的用法一致
    try:
        ax = sns.barplot(x=[s for s, _ in SCORES], y=[v for _, v in SCORES])
        ax.set_title(TEXT, fontsize=28)
        save_without_missing_glyphs(ax.figure, "seaborn.png")
        plt.close(ax.figure)
    finally:
        sns.reset_orig()


@check("OCR 繁中（單行文字讀回）")
def _():
    import pytesseract
    from matplotlib import font_manager
    from PIL import Image, ImageDraw, ImageFont
    font_path = font_manager.findfont(font_manager.FontProperties(family=["sans-serif"]))
    img = Image.new("L", (900, 160), 255)
    ImageDraw.Draw(img).text((30, 30), TEXT, fill=0, font=ImageFont.truetype(font_path, 80))
    img.save("ocr.png")
    text = compact(pytesseract.image_to_string(img, lang="chi_tra", config="--psm 7"))
    hit = sum(ch in text for ch in TEXT)
    if hit < len(TEXT) - 1:  # 容許錯一個字
        raise RuntimeError(f"只讀對 {hit}/{len(TEXT)} 字：{text[:30]}")
    return f"讀對 {hit}/{len(TEXT)} 字"


@check("Word 表格＋圖片、Excel／PowerPoint 原生圖表")
def _():
    import docx, openpyxl, pptx
    from docx.shared import Inches as DocxInches
    from openpyxl.chart import BarChart, Reference
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    d = docx.Document()
    d.add_heading(TEXT, 1)
    d.add_paragraph("國文 80 分，數學 72 分。")
    table = d.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "科目", "分數"
    for subject, score in SCORES:
        row = table.add_row().cells
        row[0].text, row[1].text = subject, str(score)
    if os.path.exists("chart.png"):
        d.add_picture("chart.png", width=DocxInches(4))
    d.save("doc.docx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["科目", "分數"])
    for row in SCORES:
        ws.append(list(row))
    chart = BarChart()
    chart.title = TEXT
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=len(SCORES) + 1), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(SCORES) + 1))
    ws.add_chart(chart, "D2")
    wb.save("sheet.xlsx")

    p = pptx.Presentation()
    slide = p.slides.add_slide(p.slide_layouts[5])  # 只有標題的版面
    slide.shapes.title.text = TEXT
    data = CategoryChartData()
    data.categories = [s for s, _ in SCORES]
    data.add_series("分數", [v for _, v in SCORES])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(8), Inches(5), data)
    p.save("slides.pptx")

    import pandas
    if pandas.read_excel("sheet.xlsx").iloc[0]["分數"] != 80:
        raise RuntimeError("Excel 讀回的內容不對")
    if len(openpyxl.load_workbook("sheet.xlsx").active._charts) != 1:
        raise RuntimeError("Excel 圖表沒有存進去")
    if len(docx.Document("doc.docx").tables[0].rows) != len(SCORES) + 1:
        raise RuntimeError("Word 表格讀回的列數不對")
    if not any(s.has_chart for s in pptx.Presentation("slides.pptx").slides[0].shapes):
        raise RuntimeError("PowerPoint 圖表沒有存進去")
    return "寫入並讀回確認"


@check("xlsxwriter 圖表、Markdown 表格（tabulate）")
def _():
    import pandas
    import xlsxwriter
    wb = xlsxwriter.Workbook("xw.xlsx")
    ws = wb.add_worksheet()
    ws.write_row(0, 0, ["科目", "分數"])
    for i, (subject, score) in enumerate(SCORES, start=1):
        ws.write_row(i, 0, [subject, score])
    chart = wb.add_chart({"type": "column"})
    chart.add_series({"categories": ["Sheet1", 1, 0, len(SCORES), 0],
                      "values": ["Sheet1", 1, 1, len(SCORES), 1], "name": TEXT})
    ws.insert_chart("D2", chart)
    wb.close()
    md = pandas.DataFrame(SCORES, columns=["科目", "分數"]).to_markdown(index=False)
    if "| 國文" not in md:
        raise RuntimeError("to_markdown 輸出不對")


@check("plotly 互動圖表（離線 HTML）")
def _():
    import plotly.express as px
    fig = px.bar(x=[s for s, _ in SCORES], y=[v for _, v in SCORES], title=TEXT)
    fig.write_html("plotly.html", include_plotlyjs=True)
    html = open("plotly.html", encoding="utf-8").read()
    if len(html) < 1_000_000:  # 內含完整 plotly.js（約 4–5 MB）才能離線打開
        raise RuntimeError("HTML 沒有內含 plotly.js，離線會打不開")
    if TEXT not in html and json.dumps(TEXT)[1:-1] not in html:  # plotly 會把中文寫成 \uXXXX
        raise RuntimeError("HTML 裡找不到圖表標題")
    return f"{len(html) // 1_000_000} MB，離線可開"


@check("graphviz 中文流程圖")
def _():
    import graphviz
    font = "Noto Sans CJK TC"
    matched = subprocess.run(["fc-match", "-f", "%{file}", font], capture_output=True, text=True).stdout
    if CJK_FILE not in matched:
        raise RuntimeError(f"系統找不到中文字型「{font}」：{matched}")
    g = graphviz.Digraph(format="png", graph_attr={"fontname": font},
                         node_attr={"fontname": font, "shape": "box"})
    g.edge("蒐集資料", TEXT)
    g.edge(TEXT, "撰寫報告")
    g.render("flow", cleanup=True)
    return f"字型：{os.path.basename(matched)}"


@check("轉成 PDF（LibreOffice）")
def _():
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf",
                    "doc.docx", "sheet.xlsx", "slides.pptx"],
                   check=True, capture_output=True, timeout=240)
    pdfs = [n for n in os.listdir() if n.endswith(".pdf")]
    if len(pdfs) < 3:
        raise RuntimeError(f"只產生 {len(pdfs)} 個 PDF（應有 3 個）")
    return "Word、Excel、PowerPoint 各轉出一個 PDF"


@check("從 PDF 讀回中文（pdfplumber、PyMuPDF、poppler）")
def _():
    import pdfplumber, pymupdf
    with pdfplumber.open("doc.pdf") as pdf:
        t1 = compact(pdf.pages[0].extract_text() or "")
    t2 = compact(pymupdf.open("doc.pdf")[0].get_text())
    t3 = compact(subprocess.run(["pdftotext", "doc.pdf", "-"], capture_output=True,
                                text=True, check=True).stdout)
    for name, t in (("pdfplumber", t1), ("PyMuPDF", t2), ("pdftotext", t3)):
        if TEXT not in t:
            raise RuntimeError(f"{name} 讀不到中文（字型可能沒嵌入）")
    subprocess.run(["pdftoppm", "-png", "-r", "50", "-f", "1", "-l", "1", "doc.pdf", "page"],
                   check=True, capture_output=True)
    return "三種方式都讀到中文，也能轉成圖片"


@check("合併列印（docxtpl＋docxcompose）")
def _():
    import docx
    from docxcompose.composer import Composer
    from docxtpl import DocxTemplate
    tpl = docx.Document()
    tpl.add_paragraph("{{ 姓名 }} 同學的國文成績：{{ 分數 }} 分")
    tpl.save("tpl.docx")
    letters = []
    for name, score in (("學生A", 80), ("學生B", 72)):
        t = DocxTemplate("tpl.docx")
        t.render({"姓名": name, "分數": score})
        t.save(f"letter_{name}.docx")
        letters.append(f"letter_{name}.docx")
    merged = Composer(docx.Document(letters[0]))
    merged.append(docx.Document(letters[1]))
    merged.save("letters.docx")
    text = "".join(p.text for p in docx.Document("letters.docx").paragraphs)
    if "學生A 同學" not in text or "學生B 同學" not in text:
        raise RuntimeError(f"合併結果不對：{text[:40]}")
    return "2 份通知單產生並合併"


@check("加密與舊格式（pypdf、pyzipper、msoffcrypto、xlrd）")
def _():
    import io
    import msoffcrypto
    import openpyxl
    import pyzipper
    import xlrd  # noqa: F401  pandas 讀 .xls 需要
    from msoffcrypto.format.ooxml import OOXMLFile
    from pypdf import PdfReader, PdfWriter
    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    w.add_blank_page(width=200, height=200)
    w.encrypt("test-pw", algorithm="AES-256")
    w.write("locked.pdf")
    r = PdfReader("locked.pdf")
    if not r.is_encrypted or not r.decrypt("test-pw") or len(r.pages) != 2:
        raise RuntimeError("PDF 加密／解密不對")
    with pyzipper.AESZipFile("locked.zip", "w", compression=pyzipper.ZIP_DEFLATED,
                             encryption=pyzipper.WZ_AES) as z:
        z.setpassword(b"test-pw")
        z.writestr("成績.txt", "國文 80")
    with pyzipper.AESZipFile("locked.zip") as z:
        z.setpassword(b"test-pw")
        if z.read("成績.txt").decode() != "國文 80":
            raise RuntimeError("加密壓縮讀回不對")
    wb = openpyxl.Workbook()
    wb.active["A1"] = "國文"
    wb.save("plain.xlsx")
    with open("plain.xlsx", "rb") as fin, open("locked.xlsx", "wb") as fout:
        OOXMLFile(fin).encrypt("test-pw", fout)
    office = msoffcrypto.OfficeFile(open("locked.xlsx", "rb"))
    office.load_key(password="test-pw")
    buf = io.BytesIO()
    office.decrypt(buf)
    if openpyxl.load_workbook(buf).active["A1"].value != "國文":
        raise RuntimeError("有密碼的 Excel 解開後內容不對")
    return "PDF、zip、Excel 加密後都能用密碼解開"


@check("掃描 PDF 轉可搜尋（ocrmypdf）")
def _():
    from PIL import Image
    if not os.path.exists("ocr.png"):
        raise RuntimeError("缺少 OCR 測試圖（上面的 OCR 項目沒有跑成功）")
    Image.open("ocr.png").convert("RGB").save("scan.pdf", resolution=150)
    run = subprocess.run(["ocrmypdf", "-l", "chi_tra", "--output-type", "pdf", "scan.pdf", "scan_ocr.pdf"],
                         capture_output=True, text=True, timeout=180)
    if run.returncode != 0:
        raise RuntimeError(f"結束碼 {run.returncode}：{run.stderr.strip()[-200:]}")
    text = compact(subprocess.run(["pdftotext", "scan_ocr.pdf", "-"], capture_output=True,
                                  text=True, check=True).stdout)
    hit = sum(ch in text for ch in TEXT)
    if hit < len(TEXT) - 1:
        raise RuntimeError(f"轉出的 PDF 只讀到 {hit}/{len(TEXT)} 字：{text[:30]}")
    return f"讀到 {hit}/{len(TEXT)} 字"


@check("公文字型對應（標楷體→楷書、新細明體→明體）")
def _():
    want = {"標楷體": "ukai", "新細明體": "uming", "微軟正黑體": "NotoSansCJK"}
    got = {}
    for family, expect in want.items():
        got[family] = subprocess.run(["fc-match", "-f", "%{file}", family],
                                     capture_output=True, text=True).stdout
        if expect not in got[family]:
            raise RuntimeError(f"{family} 對應到 {os.path.basename(got[family]) or '（無）'}")
    return "、".join(f"{k}→{os.path.basename(v)}" for k, v in got.items())


@check("簡轉繁、中文斷詞（opencc、jieba）")
def _():
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # jieba 匯入時的 pkg_resources 過時警告，無害
        import jieba
    from opencc import OpenCC
    jieba.setLogLevel(60)  # 不顯示載入字典的訊息
    converted = OpenCC("s2twp").convert("软件信息")
    if converted != "軟體資訊":
        raise RuntimeError(f"簡轉繁結果：{converted}")
    words = jieba.lcut("學生的閱讀理解需要加強")
    if len(words) < 3:
        raise RuntimeError(f"斷詞結果：{words}")
    return f"软件信息→{converted}；斷詞 {len(words)} 個"


@check("影音與壓縮檔（FFmpeg、ImageMagick、7z、unar）")
def _():
    subprocess.run(["ffmpeg", "-loglevel", "error", "-f", "lavfi", "-i", "sine=frequency=440:duration=1",
                    "tone.mp3"], check=True, capture_output=True, timeout=60)
    secs = float(subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0",
                                 "tone.mp3"], capture_output=True, text=True, check=True).stdout)
    if not 0.5 < secs < 1.5:
        raise RuntimeError(f"FFmpeg 產生的音檔長度不對：{secs}")
    from PIL import Image
    Image.new("RGB", (400, 300), "white").save("big.png")
    subprocess.run(["convert", "big.png", "-resize", "50%", "small.png"], check=True, capture_output=True)
    if Image.open("small.png").size != (200, 150):
        raise RuntimeError("ImageMagick 縮圖結果不對")
    subprocess.run(["7z", "a", "-bd", "pack.7z", "small.png"], check=True, capture_output=True)
    if not shutil.which("unar"):
        raise RuntimeError("找不到 unar")
    return f"音檔 {secs:.1f} 秒、縮圖、7z 壓縮都正常"


@check("QR code、注音、文字雲（qrcode、pypinyin、wordcloud）")
def _():
    import qrcode
    from matplotlib import font_manager
    from pypinyin import Style, pinyin
    from wordcloud import WordCloud
    qrcode.make("https://example.com/form").save("qr.png")
    zhuyin = " ".join(p[0] for p in pinyin("成績", style=Style.BOPOMOFO))
    if "ㄔㄥˊ" not in zhuyin:
        raise RuntimeError(f"注音結果：{zhuyin}")
    font = font_manager.findfont(font_manager.FontProperties(family=["sans-serif"]))
    WordCloud(font_path=font, width=300, height=200).generate("閱讀 閱讀 寫作 數學 合作").to_file("cloud.png")
    return f"成績→{zhuyin}"


@check("文字轉語音（edge-tts，線上服務）")
def _():
    import asyncio
    import edge_tts
    asyncio.run(edge_tts.Communicate(TEXT, "zh-TW-HsiaoChenNeural").save("tts.mp3"))  # 測試用文字，不含個資
    if os.path.getsize("tts.mp3") < 1000:
        raise RuntimeError("產生的語音檔太小")
    return "台灣口音語音產生成功（只用於不含個資的文字）"


@check("語音轉文字（faster-whisper，離線）")
def _():
    from faster_whisper import WhisperModel
    from opencc import OpenCC
    if not os.path.exists("tts.mp3"):
        raise RuntimeError("缺少測試音檔（上面的 edge-tts 項目沒有成功）；模型本身可能正常")
    try:
        model, device = WhisperModel("large-v3-turbo", device="cuda", compute_type="float16"), "GPU"
    except Exception:  # GPU 記憶體不足時退回 CPU
        model, device = WhisperModel("large-v3-turbo", device="cpu", compute_type="int8"), "CPU"
    segments, _info = model.transcribe("tts.mp3", language="zh")
    text = compact(OpenCC("s2twp").convert("".join(s.text for s in segments)))
    hit = sum(ch in text for ch in TEXT)
    if hit < len(TEXT) - 2:
        raise RuntimeError(f"只聽對 {hit}/{len(TEXT)} 字：{text[:30]}")
    return f"{device} 執行，聽對 {hit}/{len(TEXT)} 字"


@check("Markdown 轉 Word（pandoc）")
def _():
    with open("note.md", "w", encoding="utf-8") as f:
        f.write(f"# {TEXT}\n\n- 國文 80 分\n")
    subprocess.run(["pandoc", "note.md", "-o", "note.docx"], check=True, capture_output=True)
    import docx
    if TEXT not in docx.Document("note.docx").paragraphs[0].text:
        raise RuntimeError("轉出的 Word 內容不對")
    return ""


os.chdir("/")
shutil.rmtree(work, ignore_errors=True)

failed = [r for r in results if not r[0]]
print("🔎 文件工具自我檢查：")
for ok, name, detail in results:
    print(f"  {'✅' if ok else '❌'} {name}" + (f"（{detail}）" if detail else ""))
if failed:
    print(f"⚠️ {len(failed)} 項沒通過，AI 仍可使用，但相關功能可能要請老師換別的方式處理。")
    sys.exit(1)
print("✅ 全部通過：Word、Excel、PowerPoint、PDF、OCR、表格與圖表、合併列印、加密、公文字型、影音與語音都能用。")
